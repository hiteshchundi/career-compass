import asyncio
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import HTTPException, UploadFile

from app.ai.llm_service import AIUnavailableError, LLMService
from app.ai.docx_generator import ResumeGenerator
from app.api.routes import analyze, tailor
from app.ingestion.docx import extract_text as extract_docx_text


def resume_file() -> UploadFile:
    document = Document()
    document.add_paragraph("Bachelor of Technology. Python and Docker. 4 years experience.")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return UploadFile(file=buffer, filename="resume.docx")


def test_analyze_keeps_evidence_score_when_ai_fails(monkeypatch):
    class Unavailable:
        def summarize(self, **kwargs):
            raise AIUnavailableError("provider unavailable")

    monkeypatch.setattr(analyze, "LLMService", Unavailable)
    result = asyncio.run(analyze.analyze(resume_file(), "Requires Python, Docker, 3 years and bachelor's degree."))
    assert result["match_score"] == 100
    assert result["experience_match"] is True
    assert result["education_match"] is True
    assert result["ai_status"] == "unavailable"
    assert result["summary"] is None


def test_analyze_adds_summary_without_changing_score(monkeypatch):
    class Available:
        def summarize(self, **kwargs):
            return "Relevant experience and skills."

    monkeypatch.setattr(analyze, "LLMService", Available)
    result = asyncio.run(analyze.analyze(resume_file(), "Requires Python and Docker."))
    assert result["match_score"] == 100
    assert result["summary"] == "Relevant experience and skills."
    assert result["ai_status"] == "available"


def test_tailor_download_is_valid_docx_and_temp_file_is_removed(monkeypatch):
    class Available:
        def tailor_resume(self, **kwargs):
            return "Bachelor of Technology. Python and Docker. 4 years experience."

    monkeypatch.setattr(tailor, "LLMService", Available)
    response = asyncio.run(tailor.tailor_resume(resume_file(), "Python role"))
    path = Path(response.path)
    try:
        document = Document(path)
        assert "Python and Docker" in " ".join(p.text for p in document.paragraphs)
        assert response.filename == "tailored_resume.docx"
    finally:
        asyncio.run(response.background())
    assert not path.exists()


def test_tailor_reports_provider_failure(monkeypatch):
    class Unavailable:
        def tailor_resume(self, **kwargs):
            raise AIUnavailableError("provider unavailable")

    monkeypatch.setattr(tailor, "LLMService", Unavailable)
    try:
        asyncio.run(tailor.tailor_resume(resume_file(), "Python role"))
    except HTTPException as exc:
        assert exc.status_code == 503
        assert "temporarily unavailable" in exc.detail
    else:
        raise AssertionError("Expected a 503 response")


def test_docx_table_content_is_extracted(tmp_path):
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Bachelor of Technology, Python"
    path = tmp_path / "table-resume.docx"
    document.save(path)
    assert "Bachelor of Technology, Python" in extract_docx_text(path)


def test_invalid_resume_returns_422():
    file = UploadFile(file=BytesIO(b"not a DOCX"), filename="broken.docx")
    try:
        asyncio.run(analyze.analyze(file, "Python role"))
    except HTTPException as exc:
        assert exc.status_code == 422
    else:
        raise AssertionError("Expected an extraction error")


def test_groq_model_uses_supported_default_and_can_be_overridden(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "synthetic-key")
    monkeypatch.delenv("GROQ_MODEL", raising=False)
    assert LLMService().model == "openai/gpt-oss-120b"
    monkeypatch.setenv("GROQ_MODEL", "custom-model")
    assert LLMService().model == "custom-model"


def test_tailor_only_reorders_original_resume_sections(monkeypatch):
    service = LLMService.__new__(LLMService)
    service._complete = lambda prompt: '{"ordered_block_numbers":[0,2,1]}'
    original = "Synthetic Candidate\nSKILLS\nPython and Docker\nEXPERIENCE\n4 years experience"

    result = service.tailor_resume(original, "Python role")

    assert result.splitlines() == [
        "Synthetic Candidate",
        "EXPERIENCE",
        "4 years experience",
        "SKILLS",
        "Python and Docker",
    ]


def test_tailor_rejects_missing_or_duplicate_lines():
    service = LLMService.__new__(LLMService)
    service._complete = lambda prompt: '{"ordered_block_numbers":[0,0]}'

    try:
        service.tailor_resume("Candidate\nSKILLS\nPython", "Python role")
    except AIUnavailableError as exc:
        assert "Invalid AI tailoring response" in str(exc)
    else:
        raise AssertionError("Expected an invalid tailoring response")


def test_generated_docx_uses_headings_and_bullets_without_markdown(tmp_path):
    path = tmp_path / "tailored.docx"
    ResumeGenerator.generate(
        "# SKILLS\n- Python\n**EXPERIENCE**\n* Four years",
        path,
    )

    document = Document(path)
    paragraphs = document.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == [
        "SKILLS",
        "Python",
        "EXPERIENCE",
        "Four years",
    ]
    assert paragraphs[0].style.name.startswith("Heading")
    assert paragraphs[1].style.name == "List Bullet"
    assert all("**" not in paragraph.text for paragraph in paragraphs)
