import re

from docx import Document

from app.ai.resume_structure import clean_line, is_section_heading


class ResumeGenerator:

    @staticmethod
    def generate(resume_text: str, output_path: str):
        document = Document()

        for line in resume_text.split("\n"):
            line = clean_line(line)
            if not line or re.fullmatch(r"[-=_]{3,}", line):
                continue

            heading = line.rstrip(":").strip()
            if is_section_heading(line):
                document.add_heading(heading, level=2)
            elif re.match(r"^[*+-]\s+", line):
                document.add_paragraph(
                    re.sub(r"^[*+-]\s+", "", line),
                    style="List Bullet",
                )
            else:
                document.add_paragraph(line)

        document.save(output_path)
